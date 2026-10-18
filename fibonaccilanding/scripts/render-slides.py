#!/usr/bin/env python3
"""Готовит слайды библиотеки презентаций из pptx.

Для каждой темы из content/library/*.json, у которой есть
content/library/source/<slug>/<slug>.pptx:
  1. рендерит первые 4 слайда (LibreOffice → PDF → PNG) и сохраняет WebP
     шириной 1280px и не больше 150 КБ в public/ru/prezentacii/img/<slug>-N.webp;
  2. достаёт текст первых 4 слайдов (заголовок + тезисы) в поле slides;
  3. записывает общее число слайдов в slides_total.

Запуск:  python3 scripts/render-slides.py            — все темы с pptx
         python3 scripts/render-slides.py luna kosmos — только указанные
Нужны: libreoffice (soffice), pdftoppm (poppler-utils), python-pptx, Pillow.
"""
import glob
import io
import json
import os
import re
import subprocess
import sys
import tempfile

from PIL import Image
from pptx import Presentation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CONTENT = os.path.join(ROOT, "content", "library")
SOURCE = os.path.join(CONTENT, "source")
IMG_OUT = os.path.join(ROOT, "public", "ru", "prezentacii", "img")

PREVIEW_SLIDES = 4
WIDTH = 1280
MAX_BYTES = 150 * 1024

# служебные подписи, которые не являются содержанием слайда
NOISE = re.compile(r"^(Презентация создана в Fibonacci|PIC|\d{1,2})$")
EMU_PER_INCH = 914400


def font_size(shape):
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            if run.font.size:
                return run.font.size.pt
    return None


def clean(text):
    return re.sub(r"\s+", " ", text).strip()


def extract_slide(slide):
    """Заголовок слайда и тезисы «подзаголовок — текст»."""
    items = []
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        text = clean(sh.text_frame.text)
        if not text or NOISE.match(text):
            continue
        size = font_size(sh) or 0
        if size and size < 10:  # номер страницы, сноски
            continue
        items.append({
            "text": text,
            "size": size,
            "top": (sh.top or 0) / EMU_PER_INCH,
            "left": (sh.left or 0) / EMU_PER_INCH,
            "right": ((sh.left or 0) + (sh.width or 0)) / EMU_PER_INCH,
        })
    if not items:
        return None

    title_item = max(items, key=lambda i: i["size"])
    rest = [i for i in items if i is not title_item]
    heads = [i for i in rest if i["size"] >= 13]
    bodies = [i for i in rest if i["size"] < 13]

    def overlap(a, b):
        return min(a["right"], b["right"]) - max(a["left"], b["left"])

    pairs, used = [], set()
    for h in sorted(heads, key=lambda i: (round(i["top"], 1), i["left"])):
        below = [b for b in bodies
                 if id(b) not in used and b["top"] > h["top"] - 0.05 and overlap(h, b) > 0]
        if not below:
            below = [b for b in bodies if id(b) not in used and b["top"] > h["top"] - 0.05]
        body = min(below, key=lambda b: (b["top"] - h["top"], abs(b["left"] - h["left"]))) if below else None
        if body:
            used.add(id(body))
            pairs.append(f'{h["text"]} — {body["text"]}')
        else:
            pairs.append(h["text"])
    # тексты без подзаголовка
    for b in sorted(bodies, key=lambda i: (round(i["top"], 1), i["left"])):
        if id(b) not in used:
            pairs.append(b["text"])
    return {"title": title_item["text"], "bullets": pairs}


def render_previews(pptx_path, slug, tmp):
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "pdf", "--outdir", tmp, pptx_path],
        check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=300,
    )
    pdf = os.path.join(tmp, os.path.splitext(os.path.basename(pptx_path))[0] + ".pdf")
    prefix = os.path.join(tmp, "page")
    subprocess.run(
        ["pdftoppm", "-f", "1", "-l", str(PREVIEW_SLIDES), "-scale-to-x", str(WIDTH),
         "-scale-to-y", "-1", "-png", pdf, prefix],
        check=True, timeout=300,
    )
    pages = sorted(glob.glob(prefix + "-*.png"))
    os.makedirs(IMG_OUT, exist_ok=True)
    out = []
    for n, png in enumerate(pages, 1):
        im = Image.open(png).convert("RGB")
        if im.width != WIDTH:
            im = im.resize((WIDTH, round(im.height * WIDTH / im.width)), Image.LANCZOS)
        data = None
        for q in (82, 76, 70, 64, 58, 52, 46, 40):
            buf = io.BytesIO()
            im.save(buf, "WEBP", quality=q, method=6)
            data = buf.getvalue()
            if len(data) <= MAX_BYTES:
                break
        path = os.path.join(IMG_OUT, f"{slug}-{n}.webp")
        with open(path, "wb") as fh:
            fh.write(data)
        out.append((path, im.width, im.height, len(data), q))
    return out


def main(only):
    files = sorted(glob.glob(os.path.join(CONTENT, "*.json")))
    warnings = []
    for jf in files:
        with open(jf, encoding="utf-8") as fh:
            data = json.load(fh)
        changed = False
        for topic in data.get("topics", []):
            slug = topic["slug"]
            if only and slug not in only:
                continue
            pptx_path = os.path.join(SOURCE, slug, f"{slug}.pptx")
            if not os.path.exists(pptx_path):
                print(f"[{slug}] нет {os.path.relpath(pptx_path, ROOT)} — пропуск")
                continue
            prs = Presentation(pptx_path)
            slides = list(prs.slides)
            topic["slides_total"] = len(slides)

            with tempfile.TemporaryDirectory() as tmp:
                rendered = render_previews(pptx_path, slug, tmp)

            new_slides = []
            for n in range(1, PREVIEW_SLIDES + 1):
                old = topic["slides"][n - 1] if n - 1 < len(topic.get("slides", [])) else {}
                extracted = extract_slide(slides[n - 1]) if n - 1 < len(slides) else None
                if extracted is None:
                    warnings.append(f"[{slug}] слайд {n}: текст не найден в текстовых блоках "
                                    f"(возможно, он на картинке) — заполните title/bullets вручную")
                    extracted = {"title": old.get("title", ""), "bullets": old.get("bullets", [])}
                new_slides.append({
                    "image": f"/ru/prezentacii/img/{slug}-{n}.webp",
                    "width": rendered[n - 1][1] if n - 1 < len(rendered) else WIDTH,
                    "height": rendered[n - 1][2] if n - 1 < len(rendered) else 720,
                    "alt": old.get("alt") or f"{topic['h1']} — слайд {n}",
                    "title": extracted["title"],
                    "bullets": extracted["bullets"],
                })
            topic["slides"] = new_slides
            changed = True
            sizes = ", ".join(f"{os.path.basename(p)} {w}x{h} {b // 1024}КБ q{q}"
                              for p, w, h, b, q in rendered)
            print(f"[{slug}] слайдов всего: {len(slides)}; {sizes}")
        if changed:
            with open(jf, "w", encoding="utf-8") as fh:
                json.dump(data, fh, ensure_ascii=False, indent=2)
                fh.write("\n")
    for w in warnings:
        print("ВНИМАНИЕ:", w)


if __name__ == "__main__":
    main(set(sys.argv[1:]))
